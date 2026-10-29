"""
Bulk import of teachers and teaching allocations from Excel.

Flow for both imports:
  1. parse_*()   -> list of row dicts, each with "errors" (blocking) and "warnings"
  2. admin reviews the rows on screen (rows are kept in the session as JSON)
  3. commit_*()  -> creates records for rows without errors, in one transaction
"""
import re
from datetime import date, datetime

from django.contrib.auth.models import User
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from mainApp.models.models import Class, Stream
from mainApp.models.subject_models import Paper, Subject
from mainApp.models.teacher_models import (
    TITLE_CHOICES,
    EmploymentType,
    Qualification,
    Teacher,
    TeacherTeachingAllocation,
)
from mainApp.utils2.user_creds import (
    create_user_with_profile,
    generate_temp_password,
    generate_unique_username,
)

MAX_ROWS = 500

TEACHER_COLUMNS = [
    ("first_name", "First name *"),
    ("last_name", "Last name *"),
    ("title", "Title (Mr/Mrs/Ms/Dr ...)"),
    ("gender", "Gender (M/F)"),
    ("contact1", "Phone"),
    ("contact2", "Alternative phone"),
    ("email", "Email"),
    ("staff_no", "Staff no. (blank = auto)"),
    ("date_of_birth", "Date of birth (DD/MM/YYYY)"),
    ("date_joined", "Date joined (DD/MM/YYYY)"),
    ("employment_type", "Employment type"),
    ("qualification", "Qualification"),
    ("specialisation", "Specialisation"),
    ("registration_no", "Teacher registration no."),
    ("payroll_no", "Payroll / IPPS no."),
    ("next_of_kin_name", "Next of kin name"),
    ("next_of_kin_phone", "Next of kin phone"),
    ("next_of_kin_relationship", "Next of kin relationship"),
]

ALLOCATION_COLUMNS = [
    ("teacher", "Teacher staff no. or username *"),
    ("class_name", "Class *"),
    ("stream", "Stream (blank = whole class)"),
    ("subject_code", "Subject code *"),
    ("papers", "Paper codes, comma separated (blank = all)"),
    ("class_teacher", "Class teacher? (Yes/No)"),
]


# ---------------------------------------------------------------------------
# Small parsers
# ---------------------------------------------------------------------------

def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _phone(value):
    s = _text(value).replace(" ", "")
    # Excel drops the leading 0 of 07XXXXXXXX
    if re.fullmatch(r"7\d{8}", s):
        s = "0" + s
    return s


def _date(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = _text(value)
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"'{s}' is not a date (use DD/MM/YYYY)")


def _choice(value, choices, field_label):
    """Accept the stored value or the label, case-insensitively."""
    s = _text(value)
    if not s:
        return ""
    key = s.lower().rstrip(".")
    for stored, label in choices:
        if key in (stored.lower(), str(label).lower().rstrip("."), str(label).lower().split(" (")[0]):
            return stored
    allowed = ", ".join(str(label) for _, label in choices)
    raise ValueError(f"{field_label} '{s}' is not recognised. Use one of: {allowed}")


def _gender(value):
    s = _text(value).upper()
    if not s:
        return ""
    if s in ("M", "MALE"):
        return "M"
    if s in ("F", "FEMALE"):
        return "F"
    raise ValueError(f"Gender '{s}' should be M or F")


def _yes(value):
    return _text(value).lower() in ("yes", "y", "true", "1", "x")


def _read_rows(file_obj, columns):
    wb = load_workbook(file_obj, data_only=True, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("The file is empty.")
    # Header row is row 1; data starts at row 2. Column order follows the template.
    data = []
    for idx, raw in enumerate(rows[1:], start=2):
        if raw is None or all(v in (None, "") for v in raw):
            continue
        values = list(raw) + [None] * (len(columns) - len(raw))
        data.append((idx, {key: values[i] for i, (key, _) in enumerate(columns)}))
    if len(data) > MAX_ROWS:
        raise ValueError(f"The file has {len(data)} rows. Upload at most {MAX_ROWS} rows at a time.")
    return data


# ---------------------------------------------------------------------------
# Templates
# ---------------------------------------------------------------------------

def _template(columns, sheet_title, example_rows, validations=()):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    for col, (_, label) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col, value=label)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        ws.column_dimensions[cell.column_letter].width = max(16, len(label) + 2)
    for r, row in enumerate(example_rows, start=2):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)
    for formula, cell_range in validations:
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(cell_range)
    ws.freeze_panes = "A2"
    return wb


def teacher_template_workbook():
    example = [
        "Sarah", "Namuli", "Ms", "F", "0772123456", "", "s.namuli@example.com", "",
        "14/03/1990", "01/02/2019", "Permanent", "Degree", "Mathematics / Physics",
        "", "", "John Namuli", "0701123456", "Brother",
    ]
    emp = ",".join(label for _, label in EmploymentType.choices)
    qual = ",".join(label.split(" (")[0] for _, label in Qualification.choices)
    return _template(
        TEACHER_COLUMNS, "Teachers", [example],
        validations=[('"M,F"', "D2:D600"), (f'"{emp}"', "K2:K600"), (f'"{qual}"', "L2:L600")],
    )


def allocation_template_workbook():
    example = ["T0001", "S.1", "East", "MTC", "", "Yes"]
    return _template(
        ALLOCATION_COLUMNS, "Allocations", [example],
        validations=[('"Yes,No"', "F2:F2000")],
    )


# ---------------------------------------------------------------------------
# Teachers
# ---------------------------------------------------------------------------

def parse_teacher_file(file_obj):
    rows = []
    seen_staff_no, seen_phone, seen_email = {}, {}, {}
    existing_staff_no = set(
        s.upper() for s in Teacher.objects.exclude(staff_no__isnull=True).values_list("staff_no", flat=True)
    )
    existing_phones = set(Teacher.objects.exclude(contact1__isnull=True).values_list("contact1", flat=True))

    for line, raw in _read_rows(file_obj, TEACHER_COLUMNS):
        errors, warnings = [], []
        clean = {}

        clean["first_name"] = _text(raw["first_name"]).title()
        clean["last_name"] = _text(raw["last_name"]).title()
        if not clean["first_name"]:
            errors.append("First name is required.")
        if not clean["last_name"]:
            errors.append("Last name is required.")

        for key in ("specialisation", "registration_no", "payroll_no",
                    "next_of_kin_name", "next_of_kin_relationship"):
            clean[key] = _text(raw[key])
        clean["contact1"] = _phone(raw["contact1"])
        clean["contact2"] = _phone(raw["contact2"])
        clean["next_of_kin_phone"] = _phone(raw["next_of_kin_phone"])
        clean["staff_no"] = _text(raw["staff_no"]).upper()
        clean["email"] = _text(raw["email"]).lower()

        for key, fn in (
            ("title", lambda v: _choice(v, TITLE_CHOICES, "Title")),
            ("gender", _gender),
            ("employment_type", lambda v: _choice(v, EmploymentType.choices, "Employment type")),
            ("qualification", lambda v: _choice(v, Qualification.choices, "Qualification")),
            ("date_of_birth", _date),
            ("date_joined", _date),
        ):
            try:
                value = fn(raw[key])
                clean[key] = value.isoformat() if isinstance(value, date) else value
            except ValueError as exc:
                errors.append(str(exc))
                clean[key] = ""

        if clean["email"]:
            try:
                validate_email(clean["email"])
            except ValidationError:
                errors.append(f"Email '{clean['email']}' is not valid.")
            if clean["email"] in seen_email:
                warnings.append(f"Same email as row {seen_email[clean['email']]}.")
            seen_email.setdefault(clean["email"], line)
        else:
            warnings.append("No email: the teacher cannot verify their account by email.")

        if clean["staff_no"]:
            if clean["staff_no"] in existing_staff_no:
                errors.append(f"Staff no. {clean['staff_no']} already exists in the system.")
            if clean["staff_no"] in seen_staff_no:
                errors.append(f"Staff no. {clean['staff_no']} is repeated (row {seen_staff_no[clean['staff_no']]}).")
            seen_staff_no.setdefault(clean["staff_no"], line)

        if clean["contact1"]:
            if clean["contact1"] in existing_phones:
                warnings.append("A teacher with this phone number already exists. Check it is not a duplicate.")
            if clean["contact1"] in seen_phone:
                warnings.append(f"Same phone as row {seen_phone[clean['contact1']]}.")
            seen_phone.setdefault(clean["contact1"], line)

        rows.append({"line": line, "data": clean, "errors": errors, "warnings": warnings})
    return rows


@transaction.atomic
def commit_teacher_rows(rows, *, school):
    """
    Create teachers for rows without errors.
    Returns a list of credentials dicts (shown/downloaded once, never stored).
    """
    created = []
    for row in rows:
        if row["errors"]:
            continue
        d = row["data"]
        username = generate_unique_username(d["first_name"], d["last_name"])
        password = generate_temp_password(10)
        user, _ = create_user_with_profile(
            username=username, password=password,
            first_name=d["first_name"], last_name=d["last_name"], email=d["email"],
            school=school, role="TEACHER", allowed_levels=None,
        )
        teacher = Teacher(user=user, staff_no=d["staff_no"] or None)
        for key in ("title", "gender", "contact1", "contact2", "email", "employment_type",
                    "qualification", "specialisation", "registration_no", "payroll_no",
                    "next_of_kin_name", "next_of_kin_phone", "next_of_kin_relationship"):
            setattr(teacher, key, d.get(key) or ("" if key not in ("email", "contact1", "contact2", "gender") else None))
        teacher.date_of_birth = date.fromisoformat(d["date_of_birth"]) if d.get("date_of_birth") else None
        teacher.date_joined = (
            date.fromisoformat(d["date_joined"]) if d.get("date_joined") else timezone.localdate()
        )
        if not teacher.employment_type:
            teacher.employment_type = EmploymentType.PERMANENT
        teacher.save()
        created.append({
            "teacher_id": teacher.id, "staff_no": teacher.staff_no,
            "full_name": f"{d['first_name']} {d['last_name']}",
            "username": username, "password": password, "phone": teacher.contact1 or "",
        })
    return created


def credentials_workbook(creds):
    wb = Workbook()
    ws = wb.active
    ws.title = "Login details"
    headers = ["Staff no.", "Name", "Username", "Temporary password", "Phone"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=1, column=c, value=h).font = Font(bold=True)
        ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = 22
    for r, item in enumerate(creds, start=2):
        ws.cell(row=r, column=1, value=item["staff_no"])
        ws.cell(row=r, column=2, value=item["full_name"])
        ws.cell(row=r, column=3, value=item["username"])
        ws.cell(row=r, column=4, value=item["password"])
        ws.cell(row=r, column=5, value=item["phone"])
    return wb


# ---------------------------------------------------------------------------
# Allocations
# ---------------------------------------------------------------------------

def _find_teacher(key):
    key = _text(key)
    if not key:
        return None
    return (
        Teacher.objects.select_related("user").filter(staff_no__iexact=key).first()
        or Teacher.objects.select_related("user").filter(user__username__iexact=key).first()
    )


def _find_class(name):
    name = _text(name)
    if not name:
        return None
    return (
        Class.objects.select_related("level").filter(name__iexact=name).first()
        or Class.objects.select_related("level").filter(short_name__iexact=name).first()
    )


def parse_allocation_file(file_obj, academic_year):
    rows = []
    seen = {}
    for line, raw in _read_rows(file_obj, ALLOCATION_COLUMNS):
        errors, warnings = [], []
        out = {"teacher_id": None, "class_id": None, "stream_id": None, "subject_id": None,
               "paper_ids": [], "class_teacher": _yes(raw["class_teacher"]), "labels": {}}

        teacher = _find_teacher(raw["teacher"])
        if not teacher:
            errors.append(f"Teacher '{_text(raw['teacher'])}' not found (use staff no. or username).")
        elif teacher.status != "ACTIVE":
            errors.append(f"{teacher} is not active ({teacher.get_status_display()}).")
        else:
            out["teacher_id"] = teacher.id
            out["labels"]["teacher"] = str(teacher)

        cls = _find_class(raw["class_name"])
        if not cls:
            errors.append(f"Class '{_text(raw['class_name'])}' not found.")
        else:
            out["class_id"] = cls.id
            out["labels"]["class"] = cls.name

        stream_name = _text(raw["stream"])
        if stream_name and cls:
            stream = Stream.objects.filter(class_ref=cls, name__iexact=stream_name).first()
            if not stream:
                errors.append(f"Stream '{stream_name}' does not exist in {cls.name}.")
            else:
                out["stream_id"] = stream.id
                out["labels"]["stream"] = stream.name

        subject = Subject.objects.filter(code__iexact=_text(raw["subject_code"])).first()
        if not subject:
            errors.append(f"Subject code '{_text(raw['subject_code'])}' not found.")
        else:
            out["subject_id"] = subject.id
            out["labels"]["subject"] = f"{subject.name} ({subject.code})"
            if cls and subject.level_id != cls.level_id:
                errors.append(f"{subject.name} is a {subject.level} subject but {cls.name} is {cls.level}.")

        paper_codes = [p.strip() for p in _text(raw["papers"]).split(",") if p.strip()]
        if paper_codes and subject:
            found = list(Paper.objects.filter(subject=subject, paper_code__in=paper_codes))
            found_codes = {p.paper_code.upper() for p in found}
            missing = [c for c in paper_codes if c.upper() not in found_codes]
            if missing:
                errors.append(f"Paper(s) {', '.join(missing)} not found under {subject.code}.")
            out["paper_ids"] = [p.id for p in found]
            out["labels"]["papers"] = ", ".join(p.paper_code for p in found)

        if not errors:
            key = (out["teacher_id"], out["class_id"], out["stream_id"], out["subject_id"])
            if key in seen:
                errors.append(f"Duplicate of row {seen[key]}.")
            seen.setdefault(key, line)
            if TeacherTeachingAllocation.all_years.filter(
                teacher_id=out["teacher_id"], class_ref_id=out["class_id"], stream_id=out["stream_id"],
                subject_id=out["subject_id"], academic_year=academic_year,
            ).exists():
                warnings.append("Already allocated this year. The row will be skipped.")

        rows.append({"line": line, "data": out, "errors": errors, "warnings": warnings})
    return rows


@transaction.atomic
def commit_allocation_rows(rows, *, academic_year):
    from mainApp.services.teacher_service import assign_class_teacher

    stats = {"created": 0, "skipped": 0, "class_teachers": 0, "class_teacher_conflicts": []}
    for row in rows:
        if row["errors"]:
            continue
        d = row["data"]
        alloc, created = TeacherTeachingAllocation.all_years.get_or_create(
            teacher_id=d["teacher_id"], class_ref_id=d["class_id"], stream_id=d["stream_id"],
            subject_id=d["subject_id"], academic_year=academic_year,
        )
        if created:
            if d["paper_ids"]:
                alloc.papers.set(d["paper_ids"])
            stats["created"] += 1
        else:
            stats["skipped"] += 1

        if d["class_teacher"]:
            obj, changed = assign_class_teacher(
                class_ref=alloc.class_ref, stream=alloc.stream, teacher=alloc.teacher,
                academic_year=academic_year, replace=False,
            )
            if changed:
                stats["class_teachers"] += 1
            elif obj.teacher_id != alloc.teacher_id:
                stats["class_teacher_conflicts"].append(
                    f"{obj.class_label} already has {obj.teacher} as class teacher (row {row['line']})."
                )
    return stats
