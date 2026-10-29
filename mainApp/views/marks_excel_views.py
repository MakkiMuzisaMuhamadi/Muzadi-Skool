# mainApp/views/marks_excel_views.py
"""
Excel round-trip for marks:
  1. Teacher downloads a pre-filled sheet for one class/stream + exam type + subject/paper.
  2. Fills the Mark column (numbers or ABS) offline.
  3. Uploads it: we show a preview (new / changed / errors) before anything is saved.
  4. Confirms: marks are saved through the same rules as the Enter Marks screen,
     and every change is recorded in the change history (source = Excel upload).
"""
import io
import re
import uuid
from decimal import Decimal

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.http import HttpResponse, QueryDict
from django.shortcuts import redirect, render
from django.urls import reverse
from django.utils import timezone

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter

from mainApp.decorators import role_required
from mainApp.models.exams_models import MarkSheet
from mainApp.models.student import Student
from mainApp.marks_service import ENTRY_ROLES, MarkError, parse_mark_input, save_mark
from mainApp.views.mark_entry_views import build_session, resolve_selection
from mainApp.views.teacher_alloc_utils import get_current_school

INFO_SHEET = "_info"
DATA_SHEET = "Marks"
HEADER_ROW = 4
FIRST_DATA_ROW = 5
FILE_VERSION = "1"
MAX_UPLOAD_BYTES = 2 * 1024 * 1024
SESSION_PREFIX = "marks_import:"
SELECTION_KEYS = ("academic_year", "term", "class_ref", "stream", "exam_type", "subject", "paper")


def _selection_ids(sel):
    return {
        "academic_year": sel["academic_year"].id,
        "term": sel["term"].id,
        "class_ref": sel["class_ref"].id,
        "stream": sel["stream"].id if sel["stream"] else "",
        "exam_type": sel["exam_type"].id,
        "subject": sel["subject"].id if sel["subject"] else "",
        "paper": sel["paper"].id if sel["paper"] else "",
    }


def _as_querydict(ids: dict) -> QueryDict:
    qd = QueryDict(mutable=True)
    for k in SELECTION_KEYS:
        v = ids.get(k)
        qd[k] = "" if v in (None, "None") else str(v)
    return qd


def _target_label(sel):
    if sel["paper"]:
        return f"{sel['paper'].subject.name} - {sel['paper'].paper_name}"
    return sel["subject"].name


def _safe(text):
    return re.sub(r"[^A-Za-z0-9._-]+", "_", str(text)).strip("_")[:40]


def _excel_value_to_text(value):
    """Turn whatever Excel gives us into the text the parser expects."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        d = Decimal(str(round(float(value), 4))).normalize()
        return format(d, "f")
    return str(value).strip()


# ---------------------------------------------------------------------------
# 1. Download
# ---------------------------------------------------------------------------
@login_required
@role_required(list(ENTRY_ROLES))
def marks_excel_template(request):
    school = get_current_school(request)
    filter_form, sel = resolve_selection(request, school)
    if not sel:
        messages.error(request, "Select the class, exam type and subject/paper first.")
        return redirect("marks_entry")

    rows, _windows, _sheet = build_session(request, school, sel)
    max_mark = sel["exam_type"].max_mark

    wb = Workbook()
    ws = wb.active
    ws.title = DATA_SHEET

    stream_label = sel["stream"].name if sel["stream"] else "All streams"
    ws["A1"] = f"{sel['class_ref'].name} ({stream_label}) • {sel['exam_type'].name} • {_target_label(sel)}"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = (f"{sel['term'].name} {sel['academic_year'].year}. "
                f"Enter marks from 0 to {max_mark.normalize():f} in column D, or ABS for absent. "
                "Leave blank to keep the current mark. Do not change columns A–C.")
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A2:D2")
    ws.row_dimensions[2].height = 45

    headers = ["Candidate ID", "Student Name", "Stream", f"Mark (0–{max_mark.normalize():f} or ABS)"]
    head_fill = PatternFill("solid", fgColor="1F4E78")
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=HEADER_ROW, column=col, value=text)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center")

    input_fill = PatternFill("solid", fgColor="FFF9E6")
    r = FIRST_DATA_ROW
    for row in rows:
        ws.cell(row=r, column=1, value=row["candidate_id"])
        ws.cell(row=r, column=2, value=row["name"])
        ws.cell(row=r, column=3, value=row["stream"])
        value = row["value"]
        if value == "":
            value = None
        elif value != "ABS":
            value = float(value)
        mark_cell = ws.cell(row=r, column=4, value=value)
        if not row["locked"]:
            mark_cell.protection = Protection(locked=False)
            mark_cell.fill = input_fill
        r += 1

    for col, width in zip(range(1, 5), (18, 34, 14, 22)):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = ws.cell(row=FIRST_DATA_ROW, column=1)

    # Lock A–C so IDs cannot be changed by mistake (the upload validates anyway).
    ws.protection.sheet = True
    ws.protection.formatColumns = False

    info = wb.create_sheet(INFO_SHEET)
    info.append(["version", FILE_VERSION])
    info.append(["school", school.id])
    for k, v in _selection_ids(sel).items():
        info.append([k, v])
    info.sheet_state = "veryHidden"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = "_".join(filter(None, [
        "Marks", _safe(sel["class_ref"].name), _safe(stream_label),
        _safe(sel["exam_type"].code or sel["exam_type"].name), _safe(_target_label(sel)),
    ])) + ".xlsx"

    response = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


# ---------------------------------------------------------------------------
# 2-4. Upload, preview, confirm
# ---------------------------------------------------------------------------
def _read_workbook(upload):
    if upload.size > MAX_UPLOAD_BYTES:
        raise MarkError("The file is too large (max 2 MB).")
    if not upload.name.lower().endswith(".xlsx"):
        raise MarkError("Please upload the .xlsx file downloaded from this system.")
    try:
        wb = load_workbook(upload, data_only=True)
    except Exception:
        raise MarkError("This file could not be read. Make sure it is an Excel (.xlsx) file.")

    if INFO_SHEET not in wb.sheetnames or DATA_SHEET not in wb.sheetnames:
        raise MarkError("This is not a marks sheet from this system. Download the sheet first, then fill it in.")

    info = {}
    for key, value in wb[INFO_SHEET].iter_rows(values_only=True):
        if key:
            info[str(key)] = "" if value is None else str(value)

    rows = []
    for idx, row in enumerate(
        wb[DATA_SHEET].iter_rows(min_row=FIRST_DATA_ROW, max_col=4, values_only=True), start=FIRST_DATA_ROW
    ):
        cand, _name, _stream, mark = (list(row) + [None] * 4)[:4]
        cand = _excel_value_to_text(cand)
        if not cand:
            continue
        rows.append({"excel_row": idx, "candidate_id": cand, "raw": _excel_value_to_text(mark)})
    return info, rows


def _build_preview(request, school, sel, file_rows):
    session_rows, _w, _s = build_session(request, school, sel)
    by_cand = {r["candidate_id"]: r for r in session_rows}
    max_mark = sel["exam_type"].max_mark

    seen = set()
    preview = []
    changes = []
    for fr in file_rows:
        item = {"row": fr["excel_row"], "candidate_id": fr["candidate_id"], "name": "", "current": "",
                "new": fr["raw"], "status": "", "message": ""}
        sr = by_cand.get(fr["candidate_id"])

        if fr["candidate_id"] in seen:
            item.update(status="error", message="Candidate ID appears more than once in the file.")
        elif not sr:
            item.update(status="error", message="Not an active student in this class/stream.")
        else:
            seen.add(fr["candidate_id"])
            item["name"] = sr["name"]
            item["current"] = sr["value"]
            if fr["raw"] == "":
                item.update(status="skip", message="Blank — current mark kept.")
            elif sr["locked"]:
                item.update(status="error", message=sr["lock_reason"])
            else:
                try:
                    parsed = parse_mark_input(fr["raw"], max_mark)
                    new_display = "ABS" if parsed.is_absent else str(parsed.marks)
                    item["new"] = new_display
                    current = sr["value"]
                    same = (current == new_display) or (
                        current not in ("", "ABS") and new_display != "ABS"
                        and Decimal(current) == Decimal(new_display)
                    )
                    if same:
                        item.update(status="same", message="No change.")
                    else:
                        item.update(status="new" if current == "" else "change",
                                    message="New mark." if current == "" else f"Changes {current} → {new_display}.")
                        changes.append({"student_id": sr["id"], "marksheet_id": sr["marksheet_id"],
                                        "raw": new_display})
                except MarkError as e:
                    item.update(status="error", message=e.message)
        preview.append(item)

    counts = {k: sum(1 for p in preview if p["status"] == k) for k in ("new", "change", "same", "skip", "error")}
    return preview, changes, counts


@login_required
@role_required(list(ENTRY_ROLES))
def marks_import(request):
    school = get_current_school(request)

    # ---------- confirm ----------
    if request.method == "POST" and request.POST.get("action") == "confirm":
        token = request.POST.get("token", "")
        stored = request.session.pop(SESSION_PREFIX + token, None)
        if not stored:
            messages.error(request, "This upload has expired. Please upload the file again.")
            return redirect("marks_entry")

        qd = _as_querydict(stored["selection"])
        _form, sel = resolve_selection(request, school, data=qd)
        if not sel:
            messages.error(request, "The selection in this file is no longer valid.")
            return redirect("marks_entry")

        sheets = {s.id: s for s in MarkSheet.objects.select_related(
            "school", "exam_type", "term", "class_ref", "stream", "scheme",
        ).filter(school=school, id__in={c["marksheet_id"] for c in stored["changes"]})}
        students = {s.id: s for s in Student.objects.filter(
            school=school, id__in={c["student_id"] for c in stored["changes"]})}

        saved, failed = 0, []
        for c in stored["changes"]:
            sheet, student = sheets.get(c["marksheet_id"]), students.get(c["student_id"])
            if not sheet or not student:
                failed.append(f"Student #{c['student_id']}: not found")
                continue
            try:
                res = save_mark(user=request.user, sheet=sheet, student=student,
                                subject=sel["subject"], paper=sel["paper"], raw_value=c["raw"],
                                source="IMPORT", request=request)
                if res["result"] == "saved":
                    saved += 1
            except (MarkError, PermissionDenied) as e:
                failed.append(f"{student.last_name} {student.first_name}: {getattr(e, 'message', str(e))}")

        if saved:
            messages.success(request, f"{saved} mark(s) saved from Excel.")
        for f in failed[:15]:
            messages.error(request, f)
        if len(failed) > 15:
            messages.error(request, f"...and {len(failed) - 15} more not saved.")
        return redirect(reverse("marks_entry") + "?" + qd.urlencode())

    # ---------- upload & preview ----------
    preview = changes = counts = None
    token = ""
    sel = None

    if request.method == "POST":
        upload = request.FILES.get("file")
        if not upload:
            messages.error(request, "Choose the Excel file to upload.")
            return redirect(request.get_full_path())
        try:
            info, file_rows = _read_workbook(upload)
            if str(info.get("school")) != str(school.id):
                raise MarkError("This file was downloaded for a different school.")
            qd = _as_querydict(info)
            _form, sel = resolve_selection(request, school, data=qd)
            if not sel:
                raise MarkError("The class/exam/subject saved in this file is no longer valid.")
            if not file_rows:
                raise MarkError("No student rows were found in the file.")
            preview, changes, counts = _build_preview(request, school, sel, file_rows)
        except MarkError as e:
            messages.error(request, e.message)
            return redirect(request.get_full_path())

        # clean up older pending uploads of this user
        for key in [k for k in request.session.keys() if k.startswith(SESSION_PREFIX)]:
            del request.session[key]

        if changes:
            token = uuid.uuid4().hex
            request.session[SESSION_PREFIX + token] = {
                "selection": _selection_ids(sel),
                "changes": changes,
                "created": timezone.now().isoformat(),
            }
    else:
        _form, sel = resolve_selection(request, school)
        if not sel:
            messages.error(request, "Select the class, exam type and subject/paper first.")
            return redirect("marks_entry")

    return render(request, "dashboard/marks/marks_import.html", {
        "sel": sel,
        "target_label": _target_label(sel) if sel else "",
        "preview": preview,
        "counts": counts,
        "token": token,
        "back_query": _as_querydict(_selection_ids(sel)).urlencode() if sel else "",
    })
