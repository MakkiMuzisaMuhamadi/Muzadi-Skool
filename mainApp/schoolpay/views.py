"""
SchoolPay pages and the webhook endpoint.

    /integrations/schoolpay/                     settings + connection status
    /integrations/schoolpay/payments/            payments log + "needs attention" inbox
    /integrations/schoolpay/import-codes/        bulk import student SchoolPay codes
    /integrations/schoolpay/webhook/<token>/     called by SchoolPay's servers
"""
import csv
import io
import json
import logging

from datetime import timedelta

from django.conf import settings
from django.contrib import messages
from django.db import transaction
from django.contrib.auth.decorators import login_required
from django.core.cache import cache
from django.db.models import Q, Sum
from django.http import HttpResponse, HttpResponseBadRequest, HttpResponseForbidden, HttpResponseNotFound
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_POST

from mainApp.audit import log_action
from mainApp.decorators import role_required
from mainApp.models.fees_models import FeeCategory
from mainApp.models.schoolpay_models import SchoolPayConfig, SchoolPayFeeMapping, SchoolPayTransaction as Txn
from mainApp.models.student import Student
from mainApp.schoolpay import client, service
from mainApp.schoolpay.forms import CodeImportForm, SchoolPayConfigForm, SyncRangeForm

logger = logging.getLogger(__name__)
ADMIN_ROLES = ["SUPER_ADMIN", "ADMIN"]
# The bursar works the payments inbox; only admins change the connection or undo postings.
FEE_ROLES = ["SUPER_ADMIN", "ADMIN", "BURSAR"]


def _school(request):
    return getattr(request.user.profile, "school", None)


def _config_or_none(school):
    return SchoolPayConfig.objects.filter(school=school).first() if school else None


# ================================================================= webhook ==

def _webhook_client_ip(request):
    # Behind nginx, the real client is the LAST X-Forwarded-For entry
    # (earlier entries can be faked by the caller).
    if getattr(settings, "SCHOOLPAY_TRUST_X_FORWARDED_FOR", False):
        forwarded = request.META.get("HTTP_X_FORWARDED_FOR", "")
        if forwarded:
            return forwarded.split(",")[-1].strip()
    return request.META.get("REMOTE_ADDR", "")


@csrf_exempt
@require_POST
def schoolpay_webhook(request, token):
    config = SchoolPayConfig.objects.select_related("school").filter(webhook_token=token).first()
    if config is None or not config.is_enabled:
        return HttpResponseNotFound("Unknown or disabled endpoint")

    ip = _webhook_client_ip(request)
    if config.allowed_ip_list and ip not in config.allowed_ip_list:
        logger.warning("SchoolPay webhook from unlisted IP %s for %s", ip, config)
        return HttpResponseForbidden("IP not allowed")

    try:
        body = json.loads(request.body.decode("utf-8"))
        payment = body["payment"]
        receipt = str(payment["schoolpayReceiptNumber"])
    except (ValueError, KeyError, TypeError, UnicodeDecodeError):
        return HttpResponseBadRequest("Invalid payload")

    if not client.signature_is_valid(config.get_api_password(), receipt, body.get("signature", "")):
        logger.warning("SchoolPay webhook with bad signature (receipt %s, ip %s)", receipt, ip)
        log_action("UPDATE", model_name="SchoolPayWebhook",
                   object_repr=f"REJECTED webhook - bad signature, receipt {receipt}"[:255],
                   new_values={"ip": ip})
        return HttpResponseForbidden("Bad signature")

    txn_type = Txn.TYPE_OTHER_FEES if body.get("type") == "OTHER_FEES" else Txn.TYPE_SCHOOL_FEES
    try:
        service.ingest(config, txn_type, payment, Txn.SOURCE_WEBHOOK)
    except Exception:
        # Still answer 200: the payment is genuine and the hourly sync will
        # pick it up. A 500 would not make SchoolPay retry anyway.
        logger.exception("SchoolPay webhook: could not process receipt %s", receipt)

    SchoolPayConfig.objects.filter(pk=config.pk).update(last_webhook_at=timezone.now())
    return HttpResponse("OK")


# ================================================================ settings ==

@login_required
@role_required(ADMIN_ROLES)
def schoolpay_settings(request):
    school = _school(request)
    if not school:
        messages.error(request, "You are not assigned to any school.")
        return redirect("dashboard")

    config = _config_or_none(school) or SchoolPayConfig(school=school)
    form = SchoolPayConfigForm(request.POST or None, instance=config)

    if request.method == "POST":
        action = request.POST.get("action", "save")

        if action == "save" and form.is_valid():
            was_enabled = bool(config.pk) and SchoolPayConfig.objects.filter(pk=config.pk, is_enabled=True).exists()
            password_changed = bool(form.cleaned_data.get("api_password"))
            config = form.save(commit=False)

            # Check the code and password with SchoolPay before switching the connection on.
            if config.is_enabled and (not was_enabled or password_changed or "school_code" in form.changed_data):
                try:
                    client.fetch_day(config.school_code, config.get_api_password(), timezone.localdate())
                except client.SchoolPayError as exc:
                    config.is_enabled = False
                    config.save()
                    messages.error(request, f"Settings saved, but the connection was NOT switched on: {exc}. "
                                            f"Check the school code and API password from SchoolPay.")
                    return redirect("schoolpay_settings")
            config.save()
            log_action("UPDATE", config, new_values={
                "enabled": config.is_enabled, "strict": config.strict_mode,
                "password_changed": password_changed,
            })

            if config.is_enabled and (not was_enabled or password_changed):
                # Bring in everything paid since the start of this term (max ~3 months).
                from mainApp.fee_services import resolve_term_for_date
                from mainApp.schoolpay.tasks import sync_one_school
                today = timezone.localdate()
                term = resolve_term_for_date(school, today)
                start = max(term.start_date if term else today, today - timedelta(days=90))
                cfg_pk = config.pk
                transaction.on_commit(lambda: sync_one_school.delay(cfg_pk, start.isoformat(), today.isoformat()))
                messages.success(request, f"Connected to SchoolPay. Payments since {start:%d %b %Y} are being "
                                          f"brought in now; new payments will post automatically.")
            else:
                messages.success(request, "SchoolPay settings saved.")
            return redirect("schoolpay_settings")

        if action == "regenerate_token" and config.pk:
            config.regenerate_webhook_token()
            config.save(update_fields=["webhook_token", "updated_at"])
            messages.warning(request, "New webhook address created. Update it in your SchoolPay portal now - "
                                      "the old address no longer works.")
            return redirect("schoolpay_settings")

        if action == "test" and config.pk:
            try:
                rows = client.fetch_day(config.school_code, config.get_api_password(), timezone.localdate())
                messages.success(request, f"Connected to SchoolPay. {len(rows)} payment(s) found for today.")
            except client.SchoolPayError as exc:
                messages.error(request, f"Connection test failed: {exc}")
            return redirect("schoolpay_settings")

        if action == "save_mappings" and config.pk:
            categories = {str(c.pk): c for c in FeeCategory.objects.filter(school=school)}
            for mapping in SchoolPayFeeMapping.objects.filter(school=school):
                chosen = request.POST.get(f"mapping_{mapping.pk}", "")
                mapping.fee_category = categories.get(chosen)
                mapping.save(update_fields=["fee_category"])
            messages.success(request, "Fee mappings saved.")
            return redirect("schoolpay_settings")

    webhook_url = (request.build_absolute_uri(reverse("schoolpay_webhook", args=[config.webhook_token]))
                   if config.pk else "")
    return render(request, "dashboard/schoolpay/settings.html", {
        "title": "SchoolPay Connection",
        "form": form,
        "config": config,
        "webhook_url": webhook_url,
        "mappings": SchoolPayFeeMapping.objects.filter(school=school).select_related("fee_category"),
        "fee_categories": FeeCategory.objects.filter(school=school, is_active=True),
        "students_with_code": Student.objects.filter(school=school, is_active=True).exclude(schoolpay_code="").count(),
        "students_total": Student.objects.filter(school=school, is_active=True).count(),
        "has_current_term": bool(school.current_term_id),
        "unmapped_fees": SchoolPayFeeMapping.objects.filter(school=school, fee_category__isnull=True).count(),
        "attention_count": Txn.objects.filter(school=school, status__in=Txn.NEEDS_ATTENTION).count(),
    })


# ================================================================ payments ==

TABS = {
    "attention": ("Needs attention", Q(status__in=Txn.NEEDS_ATTENTION)),
    "pending": ("Waiting for verification", Q(status=Txn.STATUS_PENDING)),
    "posted": ("Posted", Q(status=Txn.STATUS_POSTED)),
    "ignored": ("Ignored", Q(status=Txn.STATUS_IGNORED)),
    "all": ("All", Q()),
}


def _suggest_students(school, txn):
    """Likely students for an unmatched payment, by name."""
    words = [w for w in (txn.student_name or "").split() if len(w) > 2]
    if not words:
        return []
    q = Q()
    for w in words:
        q |= Q(first_name__icontains=w) | Q(last_name__icontains=w)
    return list(Student.objects.filter(school=school, is_active=True).filter(q)
                .select_related("current_class").order_by("first_name")[:8])


@login_required
@role_required(FEE_ROLES)
def schoolpay_payments(request):
    school = _school(request)
    config = _config_or_none(school)
    if config is None:
        messages.info(request, "Set up the SchoolPay connection first.")
        return redirect("schoolpay_settings")

    base = Txn.objects.filter(school=school)
    tab = request.GET.get("tab") or ("attention" if base.filter(status__in=Txn.NEEDS_ATTENTION).exists() else "all")
    tab = tab if tab in TABS else "all"
    qs = base.filter(TABS[tab][1]).select_related("student", "fee_payment")

    search = (request.GET.get("q") or "").strip()
    if search:
        qs = qs.filter(Q(student_name__icontains=search) | Q(receipt_number__icontains=search)
                       | Q(student_payment_code__icontains=search) | Q(channel_transaction_id__icontains=search)
                       | Q(payer_detail__icontains=search))

    rows = list(qs[:300])
    for t in rows:
        t.suggestions = _suggest_students(school, t) if t.status == Txn.STATUS_UNMATCHED else []

    today = timezone.localdate()
    counts = {key: base.filter(cond).count() for key, (_label, cond) in TABS.items()}
    return render(request, "dashboard/schoolpay/payments.html", {
        "title": "SchoolPay Payments",
        "config": config,
        "rows": rows,
        "tab": tab,
        "tabs": [(k, label, counts[k]) for k, (label, _c) in TABS.items()],
        "q": search,
        "today_summary": service.reconciliation(school, today),
        "sync_form": SyncRangeForm(initial={"start": today, "end": today}),
    })


def _txn_for_action(request, pk):
    return get_object_or_404(Txn, pk=pk, school=_school(request))


def _back(request):
    from django.utils.http import url_has_allowed_host_and_scheme
    nxt = request.POST.get("next", "")
    if nxt and url_has_allowed_host_and_scheme(nxt, allowed_hosts={request.get_host()},
                                               require_https=request.is_secure()):
        return redirect(nxt)
    return redirect("schoolpay_payments")


@login_required
@role_required(FEE_ROLES)
@require_POST
def schoolpay_assign(request, pk):
    txn = _txn_for_action(request, pk)
    school = _school(request)
    student = None
    if request.POST.get("student_id"):
        student = Student.objects.filter(school=school, pk=request.POST["student_id"]).first()
    elif request.POST.get("admission_number"):
        student = Student.objects.filter(school=school,
                                         admission_number__iexact=request.POST["admission_number"].strip()).first()
    if student is None:
        messages.error(request, "Student not found. Check the admission number.")
        return _back(request)
    try:
        txn = service.assign_student(txn, student, request.user,
                                     remember_code=request.POST.get("remember_code") == "1")
        messages.success(request, f"UGX {txn.amount:,.0f} posted to {student} as receipt {txn.fee_payment.receipt_no}.")
    except service.SchoolPayActionError as exc:
        messages.error(request, str(exc))
    return _back(request)


@login_required
@role_required(FEE_ROLES)
@require_POST
def schoolpay_retry(request, pk):
    try:
        txn = service.retry(_txn_for_action(request, pk), request.user)
        messages.success(request, f"Posted as receipt {txn.fee_payment.receipt_no}.")
    except service.SchoolPayActionError as exc:
        messages.error(request, str(exc))
    return _back(request)


@login_required
@role_required(ADMIN_ROLES)
@require_POST
def schoolpay_dismiss(request, pk):
    try:
        service.dismiss(_txn_for_action(request, pk), request.user, request.POST.get("reason", ""))
        messages.success(request, "Payment dismissed.")
    except service.SchoolPayActionError as exc:
        messages.error(request, str(exc))
    return _back(request)


@login_required
@role_required(ADMIN_ROLES)
@require_POST
def schoolpay_fix_mismatch(request, pk):
    try:
        txn = service.fix_amount_mismatch(_txn_for_action(request, pk), request.user)
        messages.success(request, f"Corrected: now posted with SchoolPay's amount as receipt {txn.fee_payment.receipt_no}.")
    except service.SchoolPayActionError as exc:
        messages.error(request, str(exc))
    return _back(request)


@login_required
@role_required(FEE_ROLES)
@require_POST
def schoolpay_sync_now(request):
    school = _school(request)
    config = _config_or_none(school)
    if not config or not config.has_password:
        messages.error(request, "Save the SchoolPay code and API password first.")
        return redirect("schoolpay_settings")

    form = SyncRangeForm(request.POST)
    if request.POST.get("start") and not form.is_valid():
        messages.error(request, " ".join(form.non_field_errors()) or "Check the dates.")
        return _back(request)
    start = form.cleaned_data["start"] if form.is_valid() else None
    end = form.cleaned_data["end"] if form.is_valid() else None

    lock = f"schoolpay-sync:{school.pk}"
    if not cache.add(lock, 1, 600):
        messages.warning(request, "A sync is already running for this school. Please wait a moment.")
        return _back(request)
    try:
        s = service.sync_school(config, start, end)
        messages.success(request, f"Sync complete: {s['fetched']} checked, {s['new']} new, "
                                  f"{s['posted']} posted, {s['needs_attention']} need attention.")
    except client.SchoolPayError as exc:
        messages.error(request, f"Sync failed: {exc}")
    finally:
        cache.delete(lock)
    return _back(request)


# ============================================================ code import ==

def _read_rows(upload):
    name = upload.name.lower()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook
        ws = load_workbook(upload, read_only=True, data_only=True).active
        return [[("" if c is None else str(c)).strip() for c in row[:2]] for row in ws.iter_rows(values_only=True)]
    if name.endswith(".csv"):
        text = upload.read().decode("utf-8-sig", errors="replace")
        return [[c.strip() for c in row[:2]] for row in csv.reader(io.StringIO(text))]
    raise ValueError("Upload an .xlsx or .csv file.")


@login_required
@role_required(FEE_ROLES)
def schoolpay_import_codes(request):
    school = _school(request)
    form = CodeImportForm(request.POST or None, request.FILES or None)
    report = None

    if request.method == "POST" and form.is_valid():
        try:
            rows = _read_rows(form.cleaned_data["file"])
        except ValueError as exc:
            messages.error(request, str(exc))
            rows = None

        if rows is not None:
            report = {"updated": 0, "unchanged": 0, "not_found": [], "conflicts": []}
            for row in rows:
                if len(row) < 2 or not row[0] or not row[1]:
                    continue
                adm, code = row[0], row[1]
                if code.endswith(".0") and code[:-2].isdigit():
                    code = code[:-2]  # Excel stored the code as a decimal number
                if adm.lower().startswith("adm") and "number" in adm.lower():
                    continue  # header row
                student = Student.objects.filter(school=school, admission_number__iexact=adm).first()
                if not student:
                    report["not_found"].append(adm)
                    continue
                if student.schoolpay_code == code:
                    report["unchanged"] += 1
                    continue
                other = Student.objects.filter(school=school, schoolpay_code=code).exclude(pk=student.pk).first()
                if other:
                    report["conflicts"].append(f"{code} is already used by {other.admission_number} ({other})")
                    continue
                student.schoolpay_code = code
                student.save(update_fields=["schoolpay_code"])
                report["updated"] += 1
            log_action("UPLOAD", model_name="Student", object_repr="SchoolPay code import",
                       new_values={k: (v if isinstance(v, int) else len(v)) for k, v in report.items()})
            # Payments that were waiting for these codes can now be posted.
            config = _config_or_none(school)
            if config and report["updated"]:
                posted = 0
                for txn in Txn.objects.filter(school=school, status=Txn.STATUS_UNMATCHED):
                    try:
                        service.retry(txn, request.user)
                        posted += 1
                    except service.SchoolPayActionError:
                        pass
                report["auto_posted"] = posted

    return render(request, "dashboard/schoolpay/import_codes.html", {
        "title": "Import SchoolPay Codes", "form": form, "report": report,
    })
